"""
Social media templates for the Ninety Six Wildcats identity (v1.2).

Five post compositions (1080 x 1350) and a story (1080 x 1920), built on one
device: the number box, a slant panel shaped like the 96 box that holds the
one number that matters (a date or a score). Text in it is upright, like tags.

  announcement   date box, headline, details (maroon, or cream for academic news)
  photo          photo above a slant-edged panel holding the headline
  promo          photo, date box on the photo's edge, opponent or event title
  result         final score in two number boxes, Wildcats first
  schedule       team logo and dated rows
  story          the promo/announcement parts stacked for 9:16

Each composition returns a scene, a list of named elements, so the same layout
is drawn as a PNG (skia) and written as an editable PowerPoint file.
Run:  python social.py   -> templates/social/
"""
import io
import os
import sys

import numpy as np
import skia
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import wm  # noqa: E402
import brand as BR  # noqa: E402
import kit as K  # noqa: E402
import guide as G  # noqa: E402
import templates as T  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "templates", "social")
PHOTOS = os.path.join(HERE, "photos")
SH = BR.SPEC["shear"]
MAROON, CREAM, BLACK = K.COLORS["maroon"], K.COLORS["cream"], K.COLORS["black"]
CREAM_2 = "#EBD9C6"          # secondary text on maroon
DEEP = "#320B0B"             # photo placeholder, open dates (TV gradient bottom)
INK = "#1C1517"
W, POST_H, STORY_H = 1080, 1350, 1920
M = 72                       # outer margin
LOGO_H = 54                  # one-line team logo in the footer; two-line names get 1.6x
STORY_SAFE = 250             # app UI covers the top and bottom of stories
CAP = 0.688                  # cap height / font size, both Saira fonts


# --------------------------------------------------------------------------
# Type
# --------------------------------------------------------------------------
def display(size, color):
    return G.Style(G.COND, size, color, 0.01)


def bold(size, color):
    return G.Style(G.SAIRA, size, color, 0.0, wdth=100, wght=700)


def body(size, color):
    return G.Style(G.SAIRA, size, color, 0.0, wdth=100, wght=400)


def label(size, color):
    return G.Style(G.SAIRA, size, color, 0.14, True, wdth=100, wght=600)


_w100 = {}


def width(s, st):
    """Advance width of s in style st (shaped once at 100 px, then scaled)."""
    s = s.upper() if st.upper else s
    key = (s, st.font, st.tracking, tuple(sorted(st.axes.items())))
    if key not in _w100:
        _w100[key] = G.shaped(s, G.Style(st.font, 100, "#000", st.tracking, **st.axes))[1]
    return _w100[key] * st.size / 100


def _splits(words, n):
    if n == 1:
        yield [words]
        return
    for i in range(1, len(words) - n + 2):
        for rest in _splits(words[i:], n - 1):
            yield [words[:i]] + rest


def fit(text, max_w, max_h, max_lines=3, max_size=230, min_size=80, lh=0.9):
    """Largest display size at which the upper-cased text fits max_w x max_h
    in up to max_lines balanced lines. Returns (size, lines). Words joined by
    a no-break space (\u00a0) are never split."""
    words = text.upper().split(" ")
    best = None
    for n in range(1, min(max_lines, len(words)) + 1):
        for parts in _splits(words, n):
            lines = [" ".join(p) for p in parts]
            wmax = max(width(ln, display(100, "#000")) for ln in lines)
            size = min(max_size, max_w / wmax * 100, max_h / (CAP + (n - 1) * lh))
            score = size * (1 - 0.04 * (n - 1)) - wmax * 1e-6   # fewer lines when close, then balance
            if best is None or score > best[0]:
                best = (score, size, lines)
    return max(best[1], min_size), best[2]


def wrap(text, st, max_w):
    out, line = [], ""
    for w_ in text.split():
        trial = (line + " " + w_).strip()
        if width(trial, st) > max_w and line:
            out.append(line)
            line = w_
        else:
            line = trial
    return out + ([line] if line else [])


def block_h(size, n, lh):
    """Cap top of the first line to the baseline of the last."""
    return size * CAP + (n - 1) * size * lh


# --------------------------------------------------------------------------
# Scene elements
# --------------------------------------------------------------------------
def quad(x, y, w, h, right=False):
    """Rectangle, or a panel whose right edge leans at the logo slant."""
    p = skia.Path()
    p.moveTo(x, y)
    p.lineTo(x + w + (SH * h if right else 0), y)
    p.lineTo(x + w, y + h)
    p.lineTo(x, y + h)
    p.close()
    return p


class Scene:
    def __init__(self, h):
        self.w, self.h, self.items = W, h, []

    def shape(self, name, path, color):
        self.items.append(dict(kind="shape", name=name, path=path, color=color))

    def photo(self, name, x, y, w, h, src, focus=(0.5, 0.5), bleed=0):
        """Photo cropped to fill x, y, w, h around a focal point. bleed lets it
        run on below that box (as far as the photo goes) behind a slant edge."""
        self.items.append(dict(kind="photo", name=name, rect=(x, y, w, h), src=src, focus=focus, bleed=bleed))

    def text(self, name, lines, x, y_base, st, lh=1.0, align="left"):
        """y_base is the baseline of the first line."""
        lines = [ln.upper() if st.upper else ln for ln in lines]
        self.items.append(dict(kind="text", name=name, lines=lines, x=x, y=y_base, st=st, lh=lh, align=align))

    def logo(self, name, key, color, x, y, h):
        self.items.append(dict(kind="logo", name=name, key=key, color=color, x=x, y=y, h=h))

    def panel_text(self, name, txt, x, y, h, bg, fg, numerals=False):
        """A tag (numerals=False) or a number box (numerals=True): a slant
        panel with upright text. The number box follows the 96 box: text 64%
        of the height, 17% side padding. Returns the panel's width."""
        if numerals:
            st = G.Style(G.COND, h * BR.SPEC["box_numerals"] / CAP, fg, 0.0)
            pad = BR.SPEC["box_pad_x"] * h
        else:
            st = G.Style(G.COND, h * 0.5 / CAP, fg, 0.02)
            pad = 0.35 * h
        txt = txt.upper()
        w = width(txt, st) + 2 * pad
        self.shape(name + " panel", T.panel_path(x, y, w, h), bg)
        cap = st.size * CAP
        self.text(name, [txt], x + pad + SH * h * 0.5, y + (h + cap) / 2, st)
        return w + SH * h


def numbox_width(txt, h):
    st = G.Style(G.COND, h * BR.SPEC["box_numerals"] / CAP, "#000", 0.0)
    return width(txt.upper(), st) + 2 * BR.SPEC["box_pad_x"] * h + SH * h


def tag_width(txt, h):
    return width(txt.upper(), G.Style(G.COND, h * 0.5 / CAP, "#000", 0.02)) + 0.7 * h + SH * h


def logo_key(team):
    return "horizontal" if team in (None, "WILDCATS") else team


def logo_height(team, base=LOGO_H):
    if team in (None, "WILDCATS"):
        return base
    return base * (1.0 if len(BR.team_lines(team)[0]) == 1 else 1.6)


def footer(s, team, color, bottom):
    h = logo_height(team)
    s.logo("Team logo", logo_key(team), color, M, bottom - h, h)
    return bottom - h


def palette(d):
    """(background, headline, details, box) for the content's tone."""
    if d.get("tone") == "cream":
        return CREAM, MAROON, INK, MAROON
    return MAROON, CREAM, CREAM_2, CREAM


def title_of(d):
    if d.get("title"):
        return d["title"]
    return ("vs. " if d.get("home", True) else "at ") + d["opponent"]


# --------------------------------------------------------------------------
# Compositions
# --------------------------------------------------------------------------
def announcement(d, blank=False):
    s = Scene(POST_H)
    bg, disp, txt, boxc = palette(d)
    s.shape("Background", quad(0, 0, W, POST_H), bg)
    ftop = footer(s, d.get("team"), disp, POST_H - M)
    if blank:
        return s
    y = M
    if d.get("date"):
        s.panel_text("Date", d["date"], M, y, 190, boxc, bg, numerals=True)
        y += 190 + 70
    dst = bold(40, txt)
    det = [ln for it in d.get("details", []) for ln in wrap(it, dst, W - 2 * M)]
    dh = len(det) * 52
    kh = 34 * CAP + 44 if d.get("kicker") else 0
    size, lines = fit(d["headline"], W - 2 * M, ftop - 90 - dh - 40 - kh - y, max_lines=4)
    top = ftop - 90 - dh - 40 - block_h(size, len(lines), 0.9)      # text sits on the footer
    if d.get("kicker"):
        s.text("Kicker", [d["kicker"]], M, top - 44, label(34, disp))
    s.text("Headline", lines, M, top + size * CAP, display(size, disp), lh=0.9)
    if det:
        s.text("Details", det, M, top + block_h(size, len(lines), 0.9) + 40 + 40, dst, lh=1.3)
    return s


def photo_post(d, blank=False):
    s = Scene(POST_H)
    bg, disp, txt, _ = palette(d)
    inner = W - 2 * M - 60
    dst = bold(38, txt)
    if blank:
        size, lines, det = 120, ["", ""], []
    else:
        size, lines = fit(d["headline"], inner, 330, max_lines=3, max_size=150, min_size=84)
        det = [ln for it in d.get("details", []) for ln in wrap(it, dst, inner)]
    th = block_h(size, len(lines), 0.9)
    lh_ = logo_height(d.get("team"))
    ph = 56 + th + (24 + len(det) * 50 if det else 0) + 56 + lh_ + M
    top = POST_H - ph
    s.shape("Background", quad(0, 0, W, POST_H), DEEP)
    s.photo("Photo", 0, 0, W, top + 40, None if blank else d.get("photo"), d.get("focus", (0.5, 0.5)),
            bleed=POST_H - top - 40)
    s.shape("Panel", quad(-40, top, W - M + 40, ph + 10, right=True), bg)
    footer(s, d.get("team"), disp, POST_H - M)
    if blank:
        return s
    s.text("Headline", lines, M, top + 56 + size * CAP, display(size, disp), lh=0.9)
    if det:
        s.text("Details", det, M, top + 56 + th + 24 + 38, dst, lh=1.3)
    return s


def _date_line(d):
    parts = [d.get("day", "").upper(), (d["time"] + " PM") if d.get("time") else ""]
    return " ".join(p for p in parts if p)


def promo(d, blank=False):
    s = Scene(POST_H)
    s.shape("Background", quad(0, 0, W, POST_H), MAROON)
    bh = 210
    split = int(POST_H * 0.47)
    has_photo = blank or d.get("photo")
    if has_photo:
        s.photo("Photo", 0, 0, W, split, None if blank else d["photo"], d.get("focus", (0.5, 0.5)))
        by = split + bh * 0.42 - bh            # box top; it straddles the photo's edge
    else:
        by = M
    ftop = footer(s, d.get("team"), CREAM, POST_H - M)
    if blank:
        return s
    bw = s.panel_text("Date", d["date"], M, by, bh, CREAM, MAROON, numerals=True)
    line = _date_line(d)
    if line:
        size = min(64, (W - M - (M + bw + 30)) / width(line, display(100, CREAM)) * 100)
        s.text("Day and time", [line], M + bw + 30, by + bh, display(size, CREAM))
    y = by + bh + 70
    if d.get("kicker"):
        s.panel_text("Tag", d["kicker"], M, y, 54, CREAM, MAROON)
        y += 54 + 40
    dst = bold(40, CREAM_2)
    det = [ln for it in [d.get("place", "")] + d.get("extra", []) if it for ln in wrap(it, dst, W - 2 * M)]
    dh = len(det) * 52
    size, lines = fit(title_of(d), W - 2 * M, ftop - 64 - dh - 36 - y, max_lines=3, max_size=170)
    s.text("Headline", lines, M, y + size * CAP, display(size, CREAM), lh=0.9)
    if det:
        s.text("Details", det, M, y + block_h(size, len(lines), 0.9) + 36 + 40, dst, lh=1.3)
    return s


def result(d, blank=False):
    s = Scene(POST_H)
    s.shape("Background", quad(0, 0, W, POST_H), MAROON)
    y = M
    if d.get("photo") or blank:
        split = int(POST_H * 0.40)
        s.photo("Photo", 0, 0, W, split, None if blank else d.get("photo"), d.get("focus", (0.5, 0.5)))
        y = split + 56
    ftop = footer(s, d.get("team"), CREAM, POST_H - M)
    if blank:
        return s
    a, b = str(d["us"]), str(d["them"])
    gap = 40
    lab = bold(40, CREAM_2)
    under = 30 + 2 * 52 + (70 if d.get("note") else 0)
    head = 52 + 50
    bh = min(420, (W - 2 * M - gap) / ((numbox_width(a, 1000) + numbox_width(b, 1000)) / 1000),
             ftop - 80 - under - head - y)
    y += max(0, (ftop - 80 - y - (head + bh + under)) / 2)          # centre in the free space
    s.panel_text("Tag", "Final" + (" · Win" if d["us"] > d["them"] else ""), M, y, 52, CREAM, MAROON)
    y += head
    w1 = s.panel_text("Wildcats score", a, M, y, bh, CREAM, MAROON, numerals=True)
    x2 = M + w1 + gap
    s.panel_text("Opponent score", b, x2, y, bh, BLACK, CREAM, numerals=True)
    yb = y + bh + 30 + 40
    s.text("Wildcats label", ["WILDCATS"], M, yb, bold(40, CREAM))
    opp = wrap(d["opponent"], lab, W - M - x2)
    s.text("Opponent label", opp, x2, yb, lab, lh=1.25)
    if d.get("note"):
        s.text("Note", [d["note"]], M, yb + 52 * max(1, len(opp)) + 30, body(36, CREAM_2))
    return s


def schedule(d, blank=False):
    s = Scene(POST_H)
    s.shape("Background", quad(0, 0, W, POST_H), MAROON)
    hh = logo_height(d["team"], 76)
    s.logo("Team logo", d["team"], CREAM, M, M, hh)
    if blank:
        return s
    s.text("Title", [d["title"]], M, M + hh + 70, label(34, CREAM))
    y0 = M + hh + 130
    rows = d["events"]
    rh = min(170, (POST_H - M - 30 - y0) / len(rows))
    bh = min(rh * 0.62, 96)
    fs = 40 if rh < 125 else 46
    xcol = M + numbox_width("12/31", bh) + 34
    for k, ev in enumerate(rows):
        date, name, where, t = ev
        yt = y0 + k * rh
        off = where == "-"
        s.panel_text(f"Row {k + 1} date", date, M, yt, bh, DEEP if off else CREAM, CREAM_2 if off else MAROON,
                     numerals=True)
        name = name if off else ("vs. " if where == "H" else "at ") + name
        tst = bold(38, CREAM)
        tw = width(t, tst) if t else 0
        maxw = W - M - xcol - tw - 30
        size = fs
        while width(name, bold(size, CREAM)) > maxw and size > 30:
            size -= 1
        st = bold(size, CREAM_2 if off else CREAM)
        lines = wrap(name, st, maxw)[:2]
        mid = yt + bh / 2
        first = mid + size * CAP / 2 - (len(lines) - 1) * size * 1.05 / 2
        s.text(f"Row {k + 1} opponent", lines, xcol, first, st, lh=1.05)
        if t:
            s.text(f"Row {k + 1} time", [t], W - M, mid + 38 * CAP / 2, tst, align="right")
    return s


def story(d, blank=False):
    s = Scene(STORY_H)
    bg, disp, txt, boxc = palette(d)
    s.shape("Background", quad(0, 0, W, STORY_H), bg)
    bh = 250
    split = 1000
    if d.get("photo") or blank:
        s.photo("Photo", 0, 0, W, split, None if blank else d.get("photo"), d.get("focus", (0.5, 0.5)))
        by = split + bh * 0.42 - bh
    else:
        by = STORY_SAFE + 50
    bottom = STORY_H - STORY_SAFE - 20
    lh_ = logo_height(d.get("team"), 60)
    s.logo("Team logo", logo_key(d.get("team")), disp, M, bottom - lh_, lh_)
    if blank:
        return s
    if d.get("date"):
        bw = s.panel_text("Date", d["date"], M, by, bh, boxc, bg, numerals=True)
        line = _date_line(d)
        if line:
            size = min(76, (W - M - (M + bw + 34)) / width(line, display(100, disp)) * 100)
            s.text("Day and time", [line], M + bw + 34, by + bh, display(size, disp))
    y = by + bh + 80
    dst = bold(46, txt)
    items = [d.get("place", "")] + d.get("extra", []) + d.get("details", [])
    det = [ln for it in items if it for ln in wrap(it, dst, W - 2 * M)]
    dh = len(det) * 60
    kh = 62 + 44 if d.get("kicker") else 0
    title = d.get("headline") or title_of(d)
    size, lines = fit(title, W - 2 * M, bottom - lh_ - 80 - dh - 40 - kh - y, max_lines=3, max_size=200,
                      min_size=90)
    top = bottom - lh_ - 80 - dh - 40 - block_h(size, len(lines), 0.9)
    if d.get("kicker"):
        s.panel_text("Tag", d["kicker"], M, top - kh, 62, boxc, bg)
    s.text("Headline", lines, M, top + size * CAP, display(size, disp), lh=0.9)
    if det:
        s.text("Details", det, M, top + block_h(size, len(lines), 0.9) + 40 + 46, dst, lh=1.3)
    return s


# --------------------------------------------------------------------------
# PNG
# --------------------------------------------------------------------------
_photos = {}


def photo_image(src):
    if src not in _photos:
        im = Image.open(os.path.join(PHOTOS, src)).convert("RGBA")
        _photos[src] = skia.Image.fromarray(np.asarray(im).copy(), colorType=skia.kRGBA_8888_ColorType)
    return _photos[src]


def cover(iw, ih, w, h, focus, bleed=0):
    """Source crop (x, y, w, h) that fills w x h around a focal point, plus
    the extra height it fills below when allowed to bleed."""
    s = max(w / iw, h / ih)
    cw, ch = w / s, h / s
    sx, sy = (iw - cw) * focus[0], (ih - ch) * focus[1]
    extra = min(bleed, (ih - sy - ch) * s)
    return sx, sy, cw, ch + extra / s, h + extra


def draw(c, scene):
    for e in scene.items:
        k = e["kind"]
        if k == "shape":
            c.drawPath(e["path"], skia.Paint(AntiAlias=True, Color=wm.hexc(e["color"])))
        elif k == "photo":
            x, y, w, h = e["rect"]
            if e["src"]:
                img = photo_image(e["src"])
                sx, sy, sw, sh, h = cover(img.width(), img.height(), w, h, e["focus"], e["bleed"])
                c.drawImageRect(img, skia.Rect.MakeXYWH(sx, sy, sw, sh), skia.Rect.MakeXYWH(x, y, w, h),
                                skia.SamplingOptions(skia.FilterMode.kLinear))
            else:
                G.rect(c, x, y, w, h, DEEP)
        elif k == "text":
            for i, ln in enumerate(e["lines"]):
                G.text(c, ln, e["x"], e["y"] + i * e["st"].size * e["lh"], e["st"], align=e["align"])
        elif k == "logo":
            T.logo(c, list(_mark(e["key"])), e["color"], e["x"], e["y"], height=e["h"])


_marks = {}


def _mark(key):
    if key not in _marks:
        _marks[key] = tuple(BR.as_paths(BR.horizontal() if key == "horizontal" else BR.team(key)))
    return _marks[key]


def render_png(scene, out):
    surf = skia.Surface(scene.w, scene.h)
    draw(surf.getCanvas(), scene)
    surf.makeImageSnapshot().save(out, skia.kPNG)


# --------------------------------------------------------------------------
# PowerPoint: one slide per scene, every text element an editable text box
# --------------------------------------------------------------------------
PPT_FONTS = {G.COND: "Saira Condensed Black", G.SAIRA: "Saira"}
ASCENT = 1.135               # hhea ascender / em, both fonts
EMU = 9525                   # per px at 96 dpi


def write_pptx(slides, out, height):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * EMU), Emu(height * EMU)

    def e(v):
        return Emu(int(round(v * EMU)))

    for title, scene in slides:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        for it in scene.items:
            k = it["kind"]
            if k == "shape":
                pts = [q for cont, _ in wm.flatten(it["path"], 8.0) for q in cont]
                fb = sl.shapes.build_freeform(e(pts[0][0]), e(pts[0][1]), scale=1.0)
                fb.add_line_segments([(e(x), e(y)) for x, y in pts[1:]], close=True)
                shp = fb.convert_to_shape()
                shp.fill.solid()
                shp.fill.fore_color.rgb = RGBColor.from_string(it["color"].lstrip("#"))
                shp.line.fill.background()
                shp.name = it["name"]
            elif k == "photo":
                x, y, w, h = it["rect"]
                if it["src"]:
                    im = Image.open(os.path.join(PHOTOS, it["src"])).convert("RGB")
                    sx, sy, sw, sh, h = cover(im.width, im.height, w, h, it["focus"], it["bleed"])
                    im = im.crop((int(sx), int(sy), int(sx + sw), int(sy + sh)))
                    im.thumbnail((1600, 1600))
                else:
                    im = Image.new("RGB", (int(w) // 8, int(h) // 8), DEEP)
                buf = io.BytesIO()
                im.save(buf, "JPEG", quality=86)
                buf.seek(0)
                pic = sl.shapes.add_picture(buf, e(x), e(y), e(w), e(h))
                pic.name = it["name"] + " (right-click, Change Picture)"
            elif k == "logo":
                paths = list(_mark(it["key"]))
                b = K.bounds(paths)
                s_ = it["h"] / b.height()
                wpx, hpx = int(b.width() * s_ * 3), int(it["h"] * 3)
                surf = skia.Surface(wpx, hpx)
                cv = surf.getCanvas()
                cv.clear(skia.ColorTRANSPARENT)
                K.draw(cv, paths, it["color"], wpx / 2, hpx / 2, width=wpx)
                buf = io.BytesIO(bytes(surf.makeImageSnapshot().encodeToData()))
                pic = sl.shapes.add_picture(buf, e(it["x"]), e(it["y"]), e(b.width() * s_), e(it["h"]))
                pic.name = it["name"]
            elif k == "text":
                st = it["st"]
                n = len(it["lines"])
                pitch = st.size * it["lh"]
                widest = max(width(ln, st) for ln in it["lines"])
                bw = widest * 1.04 + 8
                x = it["x"] - (bw if it["align"] == "right" else 0)
                # box top so the first baseline lands where the PNG puts it
                top = it["y"] - ASCENT * st.size
                tb = sl.shapes.add_textbox(e(x), e(top), e(bw), e(ASCENT * st.size + (n - 1) * pitch + 0.5 * st.size))
                tb.name = it["name"]
                tf = tb.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, ln in enumerate(it["lines"]):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    if n > 1:
                        p.line_spacing = Pt(pitch * 0.75)
                    p.alignment = PP_ALIGN.RIGHT if it["align"] == "right" else PP_ALIGN.LEFT
                    r = p.add_run()
                    r.text = ln
                    f = r.font
                    f.name = PPT_FONTS[st.font]
                    f.size = Pt(st.size * 0.75)
                    f.bold = st.font == G.SAIRA and st.axes.get("wght", 400) >= 600
                    f.color.rgb = RGBColor.from_string(st.color.lstrip("#"))
                    rpr = r._r.get_or_add_rPr()
                    if st.tracking:
                        rpr.set("spc", str(int(st.tracking * st.size * 0.75 * 100)))
                    if st.upper or (st.font == G.COND and ln == ln.upper()):
                        rpr.set("cap", "all")            # typed text stays in capitals
        sl.notes_slide.notes_text_frame.text = title
    prs.save(out)


# --------------------------------------------------------------------------
# Examples (sample content; photos in source/photos)
# --------------------------------------------------------------------------
SEASON = [("8/22", "Abbeville", "H", "7:30"), ("8/29", "Westwood Christian Academy", "A", "7:00"),
          ("9/5", "Saluda", "H", "7:30"), ("9/12", "McCormick", "A", "7:30"),
          ("9/19", "Greenwood Christian School", "H", "7:30"), ("9/26", "Open date", "-", ""),
          ("10/3", "Dixie (Homecoming)", "H", "7:30"), ("10/10", "Region tournament, site TBA", "A", "TBA")]

# (composition, name, file suffix, content)
EXAMPLES = [
    ("announcement", announcement, "example", dict(
        team="WILDCATS", date="10/20", kicker="Oct 20–24", headline="Spirit Week",
        details=["Mon: Pajama Day · Tue: Decades · Wed: Jersey Day", "Thu: Twin Day · Fri: Maroon & Cream"])),
    ("announcement", announcement, "example-academic", dict(
        team="Academic Team", tone="cream", date="11/8", kicker="Academic Team", headline="Region tournament",
        details=["Saturday, Nov 8 · Matches start at 9:00 AM", "Come cheer on the team in the media center."])),
    ("photo", photo_post, "example", dict(
        team="Football", photo="football-sideline-coach.jpg", focus=(0.4, 0.45),
        headline="Game photos are up", details=["See the full gallery on the school website."])),
    ("photo", photo_post, "example-newspaper", dict(
        team="WILDCATS", tone="cream", photo="newspaper-staff.jpg", focus=(0.6, 0.4),
        headline="Meet the Weekly\u00a0Wildcat staff", details=["A new issue every Friday."])),
    ("promo", promo, "example", dict(
        team="Football", kicker="Game day", opponent="Abbeville", home=True, date="10/24", day="Friday",
        time="7:30", place="Wilson Campbell Stadium", photo="football-night.jpg", focus=(0.5, 0.32))),
    ("promo", promo, "example-band", dict(
        team="Marching Band", kicker="Competition", title="Upper State Championships", date="10/25",
        day="Saturday", time="", place="Finals performance after 8:00 PM",
        photo="band-night.jpg", focus=(0.45, 0.3))),
    ("result", result, "example", dict(
        team="Football", us=28, them=14, opponent="Abbeville", note="Region 1-2A · 6–1 overall",
        photo="football-sideline-dusk.jpg", focus=(0.5, 0.3))),
    ("result", result, "example-no-photo", dict(
        team="Basketball", us=112, them=108, opponent="McCormick", note="Final, 2OT")),
    ("schedule", schedule, "example", dict(team="Football", title="2026 Schedule", events=SEASON)),
]
STORIES = [
    ("story", story, "example", dict(
        team="WILDCATS", date="10/20", day="Monday", kicker="Spirit Week", headline="Pajama Day kicks it off",
        details=["Full list of themes on the school website."], photo="pep-rally-gym.jpg", focus=(0.5, 0.35))),
    ("story", story, "example-newspaper", dict(
        team="WILDCATS", tone="cream", kicker="Weekly Wildcat", headline="Write for us",
        details=["Staff meetings are Tuesdays after school."], photo="newspaper-staff-vertical.jpg",
        focus=(0.5, 0.45))),
]
SIZES = {"announcement": "1080 x 1350 px", "photo": "1080 x 1350 px", "promo": "1080 x 1350 px",
         "result": "1080 x 1350 px", "schedule": "1080 x 1350 px", "story": "1080 x 1920 px"}
NAMES = {"announcement": "Announcement", "photo": "Photo post", "promo": "Game or event", "result": "Result",
         "schedule": "Schedule", "story": "Story"}


def path_for(key, suffix):
    return os.path.join(OUT, f"wildcats-social-{key}-{suffix}.png")


def build():
    os.makedirs(OUT, exist_ok=True)
    for f in os.listdir(OUT):
        os.remove(os.path.join(OUT, f))
    posts, stories, done = [], [], set()
    for group, slides in ((EXAMPLES, posts), (STORIES, stories)):
        for key, fn, suffix, d in group:
            scene = fn(d)
            render_png(scene, path_for(key, suffix))
            slides.append((f"{NAMES[key]}: {suffix}", scene))
            if key not in done:
                render_png(fn(d, blank=True), path_for(key, "blank"))
                done.add(key)
            print(key, suffix)
    write_pptx(posts, os.path.join(OUT, "wildcats-social-posts.pptx"), POST_H)
    write_pptx(stories, os.path.join(OUT, "wildcats-social-stories.pptx"), STORY_H)
    print("pptx")


if __name__ == "__main__":
    build()
